"""
Tests that run WITHOUT calling the LLM (it is replaced by a fake).
Run:  uv run pytest
"""

import io

import pytest
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas

from app import main, pipeline
from app.csv_export import rows_to_csv
from app.extractor import ExtractionError
from app.pipeline import process_file
from app.schemas import Receipt, ReceiptRow


def make_pdf(text: str = "Coffee Shop total 5.00") -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    if text:
        pdf.drawString(50, 750, text)
    pdf.save()
    return buffer.getvalue()


GOOD = Receipt(date="2026-01-02", vendor="Cafe", amount=5.0, tax=0.4, currency="USD", category="Meals")


# ---------- Schema cleaning ----------
def test_receipt_cleans_messy_values():
    r = Receipt(date="2026-02-30", vendor="  ", amount="$1,234.50", tax="12,50", currency="usd", category="meals")
    assert r.date is None  # not a real date
    assert r.vendor is None
    assert r.amount == 1234.50
    assert r.tax == 12.50
    assert r.currency == "USD"
    assert r.category == "Meals"


def test_receipt_rejects_unknown_category_and_negative_amount():
    r = Receipt(amount=-5, category="Pizza", currency="DOLLARS")
    assert r.amount is None and r.category is None and r.currency is None


# ---------- Pipeline ----------
def test_complete_receipt_is_ok(monkeypatch):
    monkeypatch.setattr(pipeline, "extract_receipt", lambda **kw: GOOD)
    row = process_file("a.pdf", make_pdf())
    assert row.status == "ok" and row.notes == ""


def test_missing_field_is_flagged_not_guessed(monkeypatch):
    monkeypatch.setattr(pipeline, "extract_receipt", lambda **kw: Receipt(vendor="Cafe", amount=5.0))
    row = process_file("a.pdf", make_pdf())
    assert row.status == "needs_review"
    assert row.date is None
    assert "date" in row.notes and "category" in row.notes


def test_tax_missing_alone_is_fine(monkeypatch):
    monkeypatch.setattr(pipeline, "extract_receipt", lambda **kw: GOOD.model_copy(update={"tax": None}))
    assert process_file("a.pdf", make_pdf()).status == "ok"


def test_tax_bigger_than_total_is_flagged(monkeypatch):
    monkeypatch.setattr(pipeline, "extract_receipt", lambda **kw: GOOD.model_copy(update={"tax": 99}))
    assert process_file("a.pdf", make_pdf()).status == "needs_review"


def test_llm_failure_becomes_error_row(monkeypatch):
    def boom(**kw):
        raise ExtractionError("rate limit")

    monkeypatch.setattr(pipeline, "extract_receipt", boom)
    row = process_file("a.pdf", make_pdf())
    assert row.status == "error" and "rate limit" in row.notes


def test_unexpected_bug_does_not_crash(monkeypatch):
    def boom(**kw):
        raise RuntimeError("bug")

    monkeypatch.setattr(pipeline, "extract_receipt", boom)
    assert process_file("a.pdf", make_pdf()).status == "error"


@pytest.mark.parametrize(
    "name,data,message",
    [
        ("notes.txt", b"hello", "Unsupported"),
        ("empty.pdf", b"", "empty"),
        ("broken.pdf", b"this is not a pdf", "Could not read"),
        ("blank.pdf", make_pdf(""), "no readable text"),
        ("huge.png", b"x" * (11 * 1024 * 1024), "larger"),
        ("", b"abc", "Unsupported"),
    ],
)
def test_bad_files_give_clear_errors(name, data, message):
    row = process_file(name, data)
    assert row.status == "error" and message in row.notes


# ---------- CSV ----------
def test_csv_blocks_formula_injection_and_handles_none():
    row = ReceiptRow(file="a.pdf", status="ok", vendor="=HYPERLINK(1)", amount=5.0)
    lines = rows_to_csv([row]).splitlines()
    assert lines[0] == "file,date,vendor,amount,tax,currency,category,status,notes"
    assert "'=HYPERLINK(1)" in lines[1]
    assert lines[1].startswith("a.pdf,,")  # None -> empty cell


# ---------- API ----------
@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(main, "GROQ_API_KEY", "test-key")
    monkeypatch.setattr(pipeline, "extract_receipt", lambda **kw: GOOD)
    return TestClient(main.app)


def test_home_and_health(client):
    assert client.get("/").status_code == 200
    assert client.get("/health").json()["status"] == "ok"


def test_extract_batch_mixes_good_and_bad(client):
    files = [("files", ("a.pdf", make_pdf(), "application/pdf")), ("files", ("b.txt", b"x", "text/plain"))]
    body = client.post("/extract", files=files).json()
    assert (body["total"], body["ok"], body["errors"]) == (2, 1, 1)
    assert body["rows"][1]["status"] == "error"


def test_extract_csv_download(client):
    response = client.post("/extract/csv", files=[("files", ("a.pdf", make_pdf(), "application/pdf"))])
    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]


def test_no_files_and_too_many_files(client):
    assert client.post("/extract", files=[]).status_code in (400, 422)
    many = [("files", (f"{i}.pdf", make_pdf(), "application/pdf")) for i in range(21)]
    assert client.post("/extract", files=many).status_code == 400


def test_missing_api_key_gives_clear_error(monkeypatch):
    monkeypatch.setattr(main, "GROQ_API_KEY", None)
    response = TestClient(main.app).post("/extract", files=[("files", ("a.pdf", make_pdf(), "application/pdf"))])
    assert response.status_code == 500 and "GROQ_API_KEY" in response.json()["detail"]


def test_xlsx_has_real_dates_widths_and_header(client):
    from datetime import date

    from openpyxl import load_workbook

    body = client.post("/extract", files=[("files", ("a.pdf", make_pdf(), "application/pdf"))]).json()
    # The web page sends the rows it already has to /export/xlsx (no 2nd LLM call)
    exported = client.post("/export/xlsx", json={"rows": body["rows"]})
    assert exported.status_code == 200
    workbook = load_workbook(io.BytesIO(exported.content))
    assert workbook.sheetnames == ["Expense Report"]  # everything on one sheet
    sheet = workbook.active

    # Summary part
    assert sheet["B9"].value == "Meals" and sheet["E9"].value == 5.0
    assert sheet["B10"].value == "Total USD"
    assert sheet["J5"].value == 5.0  # "Total spent" card
    pie_data = sheet._charts[0].series[0].val.numRef.f
    assert "$E$9" in pie_data  # pie uses the Total column (not Tax)

    # Receipts table part (below the summary): find its header row
    header_row = next(r for r in range(1, 60) if sheet[f"C{r}"].value == "Date")
    assert sheet[f"B{header_row}"].value == "File"
    assert sheet[f"C{header_row + 1}"].value.date() == date(2026, 1, 2)  # a real date, not text
    assert sheet.column_dimensions["C"].width >= 12  # wide enough: no #####
    response = client.post("/extract/xlsx", files=[("files", ("a.pdf", make_pdf(), "application/pdf"))])
    assert response.content[:2] == b"PK"  # xlsx files are zip files


def test_settings_and_static_files(client):
    settings = client.get("/settings").json()
    assert settings["max_files"] == 20 and ".pdf" in settings["allowed_extensions"]
    assert client.get("/static/app.js").status_code == 200
    assert client.get("/static/styles.css").status_code == 200


def test_export_csv_and_validation(client):
    rows = [{"file": "a.pdf", "status": "ok", "vendor": "Cafe", "amount": "12.50", "date": "2026-13-45"}]
    response = client.post("/export/csv", json={"rows": rows})
    assert response.status_code == 200
    assert "a.pdf,,Cafe,12.5" in response.text  # bad date dropped, amount cleaned by Pydantic
    assert client.post("/export/xlsx", json={"rows": []}).status_code == 400
    assert client.post("/export/csv", json={"wrong": 1}).status_code == 422
